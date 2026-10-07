"""Upload validation and Unstructured parsing behavior."""

from io import BytesIO
from types import SimpleNamespace

import pytest
from pydantic import ValidationError

from document_rag.config import AppConfig
from document_rag.models import UploadedDocument
from document_rag.parsing import (
    DocumentParsingError,
    UploadValidationError,
    parse_documents,
    validate_uploads,
)


def upload(name: str, content: bytes) -> UploadedDocument:
    return UploadedDocument(filename=name, content=content)


def test_validation_rejects_empty_unsupported_and_invalid_text() -> None:
    with pytest.raises(ValidationError, match="at least 1 byte"):
        upload("empty.txt", b"")
    with pytest.raises(ValidationError, match="Only PDF, TXT, Markdown, DOC, DOCX, and CSV"):
        upload("picture.png", b"image")
    with pytest.raises(UploadValidationError, match="contains no text"):
        validate_uploads([upload("blank.md", b"  \r\n  ")])
    with pytest.raises(UploadValidationError, match="UTF-8"):
        validate_uploads([upload("bad.txt", b"\xff")])


def test_validation_rejects_set_limits_and_duplicate_names() -> None:
    with pytest.raises(UploadValidationError, match="at least one"):
        validate_uploads([])
    with pytest.raises(UploadValidationError, match="at most 1"):
        validate_uploads(
            [upload("one.txt", b"one"), upload("two.md", b"two")],
            AppConfig(max_files=1),
        )
    with pytest.raises(UploadValidationError, match="exceeds"):
        validate_uploads([upload("long.txt", b"long")], AppConfig(max_file_bytes=3))
    with pytest.raises(UploadValidationError, match="Duplicate filename"):
        validate_uploads([upload("Notes.md", b"one"), upload("notes.MD", b"two")])


def test_real_text_and_markdown_are_separate_documents() -> None:
    parsed = parse_documents(
        [
            upload("notes.txt", b"This is a plain text note with one fact."),
            upload(
                "prices.md",
                b"# Pricing\n\n| Item | Price |\n| --- | --- |\n| Apple | 2 |",
            ),
        ]
    )

    assert [document.source_name for document in parsed] == ["notes.txt", "prices.md"]
    assert parsed[0].source_id != parsed[1].source_id
    assert all(element.source_name == "notes.txt" for element in parsed[0].elements)
    assert all(element.source_name == "prices.md" for element in parsed[1].elements)
    assert any("plain text note" in element.text for element in parsed[0].elements)
    table = next(element for element in parsed[1].elements if element.category == "Table")
    assert table.section_title == "Pricing"
    assert table.table_headers == ["Item", "Price"]
    assert "Apple" in table.text and "2" in table.text
    assert "<table" in table.table_html


@pytest.mark.parametrize(
    "content",
    [b"Item,Quantity\nPens,12\nPencils,7\n", b"\xef\xbb\xbfItem,Quantity\nPens,12\nPencils,7\n"],
)
def test_real_csv_keeps_header_and_rows(content: bytes) -> None:
    parsed = parse_documents([upload("inventory.csv", content)])

    assert parsed[0].source_name == "inventory.csv"
    assert len(parsed[0].elements) == 1
    table = parsed[0].elements[0]
    assert table.category == "Table"
    assert table.table_headers == ["Item", "Quantity"]
    assert all(value in table.text for value in ("Item", "Quantity", "Pens", "12"))
    assert "<table" in table.table_html


def test_real_docx_keeps_title_and_table_structure() -> None:
    from docx import Document

    document = Document()
    document.add_heading("Quarterly results", level=1)
    document.add_paragraph("The total is shown below.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Value"
    row = table.add_row().cells
    row[0].text = "Sales"
    row[1].text = "42"
    buffer = BytesIO()
    document.save(buffer)

    parsed = parse_documents([upload("results.docx", buffer.getvalue())])

    assert parsed[0].source_name == "results.docx"
    assert any(element.category == "Title" for element in parsed[0].elements)
    parsed_table = next(element for element in parsed[0].elements if element.category == "Table")
    assert parsed_table.section_title == "Quarterly results"
    assert all(value in parsed_table.text for value in ("Item", "Value", "Sales", "42"))
    assert "<table" in parsed_table.table_html


def test_legacy_doc_is_dispatched_to_unstructured(monkeypatch) -> None:
    from document_rag import parsing

    calls = []

    def fake_partition(**kwargs):
        calls.append(kwargs)
        return [SimpleNamespace(text="Legacy content", category="NarrativeText", metadata=None)]

    monkeypatch.setattr("unstructured.partition.auto.partition", fake_partition)
    result = parsing.parse_documents([upload("legacy.doc", b"legacy binary")])

    assert result[0].elements[0].text == "Legacy content"
    assert calls[0]["content_type"] == "application/msword"
    assert calls[0]["infer_table_structure"] is True


def test_pdf_metadata_table_and_section_are_preserved(monkeypatch: pytest.MonkeyPatch) -> None:
    import unstructured.partition.auto as auto

    seen: dict[str, object] = {}

    def fake_partition(**kwargs: object) -> list[SimpleNamespace]:
        seen.update(kwargs)
        return [
            SimpleNamespace(
                category="Title",
                text="Annual results",
                metadata=SimpleNamespace(page_number=2, text_as_html=None),
            ),
            SimpleNamespace(
                category="Table",
                text="Year Revenue 2026 42",
                metadata=SimpleNamespace(
                    page_number=3,
                    text_as_html=(
                        "<table><thead><tr><th>Year</th><th>Revenue</th></tr></thead>"
                        "<tbody><tr><td>2026</td><td>42</td></tr></tbody></table>"
                    ),
                ),
            ),
        ]

    monkeypatch.setattr(auto, "partition", fake_partition)
    parsed = parse_documents([upload("report.pdf", b"%PDF-fake-for-dispatch-test")])

    assert seen["content_type"] == "application/pdf"
    assert seen["metadata_filename"] == "report.pdf"
    assert seen["strategy"] == "hi_res"
    assert seen["pdf_infer_table_structure"] is True
    assert isinstance(seen["file"], BytesIO)
    table = parsed[0].elements[1]
    assert table.category == "Table"
    assert table.page_number == 3
    assert table.section_title == "Annual results"
    assert table.table_headers == ["Year", "Revenue"]
    assert "2026 42" in table.text
    assert "<thead>" in table.table_html


def test_file_with_no_extracted_text_is_rejected(monkeypatch: pytest.MonkeyPatch) -> None:
    import unstructured.partition.auto as auto

    monkeypatch.setattr(auto, "partition", lambda **kwargs: [])
    with pytest.raises(DocumentParsingError, match="No extractable text"):
        parse_documents([upload("scan.pdf", b"%PDF-fake")])


def test_pdf_falls_back_to_fast_when_tesseract_is_missing(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from unstructured_pytesseract.pytesseract import TesseractNotFoundError
    import unstructured.partition.auto as auto

    strategies = []

    def fake_partition(**kwargs):
        strategies.append((kwargs["strategy"], kwargs["pdf_infer_table_structure"]))
        if kwargs["strategy"] == "hi_res":
            raise TesseractNotFoundError()
        return [
            SimpleNamespace(
                category="NarrativeText",
                text="Employees receive 20 days of leave.",
                metadata=SimpleNamespace(page_number=1),
            )
        ]

    monkeypatch.setattr(auto, "partition", fake_partition)
    parsed = parse_documents([upload("policy.pdf", b"%PDF-fake")])

    assert strategies == [("hi_res", True), ("fast", False)]
    assert parsed[0].elements[0].page_number == 1
    assert "20 days" in parsed[0].elements[0].text
    assert "table structure" in parsed[0].warnings[0]


def test_pdf_without_extractable_text_gives_tesseract_setup_hint(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from unstructured_pytesseract.pytesseract import TesseractNotFoundError
    import unstructured.partition.auto as auto

    def fake_partition(**kwargs):
        if kwargs["strategy"] == "hi_res":
            raise TesseractNotFoundError()
        return []

    monkeypatch.setattr(auto, "partition", fake_partition)
    with pytest.raises(DocumentParsingError, match="Install Tesseract OCR"):
        parse_documents([upload("scan.pdf", b"%PDF-fake")])


def test_real_text_pdf_can_be_partitioned_without_tesseract() -> None:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    from document_rag.parsing import _partition_upload

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 14 Tf 72 720 Td (Employees receive 20 days of annual leave.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = BytesIO()
    writer.write(buffer)

    elements = _partition_upload(upload("policy.pdf", buffer.getvalue()), pdf_strategy="fast")

    assert any("20 days of annual leave" in element.text for element in elements)


def test_hosted_unstructured_api_is_used_when_key_is_configured(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import json
    from pathlib import Path
    import unstructured_client

    seen = {}

    def fake_client(**kwargs):
        seen["key"] = kwargs["api_key_auth"]
        seen["server_url"] = kwargs["server_url"]

        def partition(*, request):
            parameters = request.partition_parameters
            seen["strategy"] = parameters.strategy
            seen["filename"] = parameters.files.file_name
            seen["table_structure"] = parameters.pdf_infer_table_structure
            return SimpleNamespace(
                status_code=200,
                raw_response=SimpleNamespace(text=json.dumps([
                    {
                        "type": "Title",
                        "element_id": "title-1",
                        "text": "Leave",
                        "metadata": {"page_number": 2},
                    },
                    {
                        "type": "NarrativeText",
                        "element_id": "text-1",
                        "text": "Employees receive 20 days of leave.",
                        "metadata": {"page_number": 2},
                    },
                    {
                        "type": "Table",
                        "text": "Type Days Annual 20",
                        "metadata": {
                            "page_number": 2,
                            "text_as_html": "<table><tr><th>Type</th><th>Days</th></tr><tr><td>Annual</td><td>20</td></tr></table>",
                        },
                    },
                ]))
            )

        return SimpleNamespace(general=SimpleNamespace(partition=partition))

    monkeypatch.setattr(unstructured_client, "UnstructuredClient", fake_client)
    parsed = parse_documents(
        [upload("policy.pdf", b"%PDF-fake")],
        config=AppConfig(
            unstructured_api_key="test-partition-key",
            unstructured_api_url="https://example.api.unstructuredapp.io/general/v0/general",
        ),
    )

    assert seen["key"] == "test-partition-key"
    assert seen["server_url"] == "https://example.api.unstructuredapp.io"
    assert Path(seen["filename"]).name == "policy.pdf"
    assert not Path(seen["filename"]).exists()
    assert seen["table_structure"] is True
    assert parsed[0].elements[1].section_title == "Leave"
    assert parsed[0].elements[1].page_number == 2
    assert "20 days" in parsed[0].elements[1].text
    assert parsed[0].elements[2].table_headers == ["Type", "Days"]
    assert parsed[0].elements[2].source_name == "policy.pdf"


def test_key_without_partition_url_keeps_local_parsing(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import document_rag.parsing as parsing
    import unstructured.partition.auto as auto

    monkeypatch.setattr(
        parsing,
        "_partition_via_api",
        lambda *args: (_ for _ in ()).throw(AssertionError("Hosted API called")),
    )
    monkeypatch.setattr(
        auto,
        "partition",
        lambda **kwargs: [
            SimpleNamespace(text="Local text", category="NarrativeText", metadata=None)
        ],
    )

    parsed = parse_documents(
        [upload("notes.txt", b"Local text")],
        config=AppConfig(unstructured_api_key="unmatched-key"),
    )
    assert parsed[0].elements[0].text == "Local text"


@pytest.mark.parametrize("failure", ["connection", "timeout"])
def test_hosted_errors_give_specific_setup_guidance(monkeypatch, failure) -> None:
    import httpx
    import document_rag.parsing as parsing

    error, message = {
        "connection": (httpx.ConnectError("private diagnostics"), "DNS, proxy, and TLS"),
        "timeout": (httpx.ReadTimeout("private diagnostics"), "timed out"),
    }[failure]

    def fail(*args):
        raise error

    monkeypatch.setattr(parsing, "_partition_via_api", fail)
    with pytest.raises(DocumentParsingError, match=message) as caught:
        parse_documents(
            [upload("notes.txt", b"Test")],
            config=AppConfig(
                unstructured_api_key="test-key",
                unstructured_api_url="https://api.unstructuredapp.io",
            ),
        )
    assert "private diagnostics" not in str(caught.value)


def test_unauthorized_hosted_parsing_falls_back_locally_for_all_files(monkeypatch) -> None:
    import document_rag.parsing as parsing

    class UnauthorizedError(Exception):
        status_code = 401

    calls = []

    def reject(*args):
        calls.append("api")
        raise UnauthorizedError("private diagnostics")

    def local(upload, *, pdf_strategy="hi_res"):
        calls.append(upload.filename)
        return [SimpleNamespace(text="Local content", category="NarrativeText", metadata=None)]

    monkeypatch.setattr(parsing, "_partition_via_api", reject)
    monkeypatch.setattr(parsing, "_partition_upload", local)
    parsed = parse_documents(
        [upload("first.txt", b"First"), upload("second.md", b"Second")],
        config=AppConfig(
            unstructured_api_key="test-key",
            unstructured_api_url="https://api.unstructuredapp.io",
        ),
    )
    assert calls == ["api", "first.txt", "second.md"]
    assert all(doc.elements[0].text == "Local content" for doc in parsed)
    assert all("parsed locally" in doc.warnings[0] for doc in parsed)
    assert all("private diagnostics" not in doc.warnings[0] for doc in parsed)
